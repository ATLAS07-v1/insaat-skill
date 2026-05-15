#!/usr/bin/env python3
"""Extract technical specification clause candidates from common document exports."""

from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable


NORMATIVE_KEYWORDS = [
    "shall",
    "must",
    "required",
    "requirements",
    "compliance",
    "olmalıdır",
    "zorunlu",
    "gereklidir",
    "gerekir",
    "yapılacaktır",
    "sağlanacaktır",
    "uygulanacaktır",
    "teslim edilecektir",
    "onaylanacaktır",
]

EVIDENCE_KEYWORDS = {
    "test_report": ["test", "deney", "laboratuvar", "rapor", "numune"],
    "certificate": ["sertifika", "certificate", "uygunluk belgesi", "ce", "tse"],
    "submittal": ["submittal", "onay", "teknik föy", "datasheet", "malzeme onayı"],
    "photo": ["foto", "fotoğraf", "photo", "görsel"],
    "measurement": ["ölçüm", "tolerans", "mm", "cm", "mpa", "%", "±"],
    "method_statement": ["method statement", "uygulama metodu", "prosedür"],
    "as_built": ["as-built", "as built", "iş sonu", "o&m", "bakım"],
    "ids_report": ["ifc", "ids", "property", "classification", "parameter"],
}

CATEGORY_KEYWORDS = {
    "material": ["malzeme", "ürün", "datasheet", "teknik föy", "sertifika", "material"],
    "workmanship": ["uygulama", "montaj", "serim", "döküm", "kür", "derz", "işçilik"],
    "test": ["test", "deney", "rapor", "laboratuvar", "numune"],
    "tolerance": ["tolerans", "limit", "minimum", "maximum", "en az", "en çok", "mm", "±"],
    "submittal": ["submittal", "onay", "teslim", "as-built", "o&m", "garanti"],
    "bim": ["ifc", "ids", "property", "classification", "parameter", "bim"],
}

CRITICAL_KEYWORDS = [
    "yangın",
    "fire",
    "taşıyıcı",
    "structural",
    "statik",
    "güvenlik",
    "safety",
    "su yalıtımı",
    "waterproofing",
    "zorunlu",
    "must",
    "shall",
]

CLAUSE_PREFIX_RE = re.compile(r"^\s*((?:\d{1,2}(?:\.\d{1,3})+)|(?:\d{2}\s+\d{2}\s+\d{2})|(?:[A-Z]\.)|(?:[a-z]\))|(?:\d+\)))\s*(.+)")


def read_text(path: Path) -> tuple[str, list[str]]:
    suffix = path.suffix.lower()
    warnings: list[str] = []
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace"), warnings
    if suffix == ".csv":
        rows: list[str] = []
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            for row in csv.reader(handle):
                rows.append(" | ".join(cell.strip() for cell in row if cell.strip()))
        return "\n".join(rows), warnings
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        return json.dumps(data, ensure_ascii=False, indent=2), warnings
    if suffix == ".pdf":
        text = try_pypdf(path, warnings)
        if not text.strip():
            text = try_pdfplumber(path, warnings)
        if not text.strip():
            warnings.append("PDF metni çıkarılamadı; taranmış PDF veya eksik PDF aracı olabilir.")
        return text, warnings
    if suffix == ".docx":
        return read_docx(path, warnings), warnings
    raise SystemExit(f"Desteklenmeyen dosya türü: {suffix}")


def try_pypdf(path: Path, warnings: list[str]) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        warnings.append("pypdf modülü yok; PDF için pypdf atlandı.")
        return ""
    try:
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # pragma: no cover - depends on external PDFs
        warnings.append(f"pypdf PDF okuma hatası: {exc}")
        return ""


def try_pdfplumber(path: Path, warnings: list[str]) -> str:
    try:
        import pdfplumber  # type: ignore
    except Exception:
        warnings.append("pdfplumber modülü yok; PDF tablo/metin çıkarımı atlandı.")
        return ""
    try:
        chunks: list[str] = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                chunks.append(page.extract_text() or "")
                for table in page.extract_tables() or []:
                    for row in table:
                        chunks.append(" | ".join(cell or "" for cell in row))
        return "\n".join(chunks)
    except Exception as exc:  # pragma: no cover - depends on external PDFs
        warnings.append(f"pdfplumber PDF okuma hatası: {exc}")
        return ""


def read_docx(path: Path, warnings: list[str]) -> str:
    try:
        from docx import Document  # type: ignore
    except Exception:
        warnings.append("python-docx modülü yok; DOCX okunamadı.")
        return ""
    try:
        document = Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
        return "\n".join(parts)
    except Exception as exc:  # pragma: no cover - depends on external docs
        warnings.append(f"DOCX okuma hatası: {exc}")
        return ""


def normalize_lines(text: str) -> list[str]:
    lines = []
    for raw in text.splitlines():
        line = re.sub(r"\s+", " ", raw).strip()
        if line:
            lines.append(line)
    return lines


def contains_any(text: str, keywords: Iterable[str]) -> bool:
    lower = text.lower()
    return any(keyword.lower() in lower for keyword in keywords)


def infer_category(text: str) -> str:
    lower = text.lower()
    for category, keywords in CATEGORY_KEYWORDS.items():
        if any(keyword.lower() in lower for keyword in keywords):
            return category
    return "general"


def infer_evidence(text: str) -> list[str]:
    lower = text.lower()
    required = [
        evidence_type
        for evidence_type, keywords in EVIDENCE_KEYWORDS.items()
        if any(keyword.lower() in lower for keyword in keywords)
    ]
    return required or ["inspection_record"]


def infer_severity(text: str) -> str:
    if contains_any(text, CRITICAL_KEYWORDS):
        return "high"
    if contains_any(text, ["should", "önerilir", "may", "mümkünse"]):
        return "low"
    return "medium"


def extract_acceptance_criteria(text: str) -> str | None:
    patterns = [
        r"(?:minimum|min\.?|en az)\s+[^.;,]+",
        r"(?:maximum|max\.?|en çok)\s+[^.;,]+",
        r"(?:tolerans|tolerance|limit)\s*[:=]?\s*[^.;,]+",
        r"±\s*\d+(?:[.,]\d+)?\s*(?:mm|cm|m|%)?",
        r"\d+(?:[.,]\d+)?\s*(?:mm|cm|m2|m3|mpa|%)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return match.group(0).strip()
    return None


def extract_clauses(lines: list[str], min_length: int, include_all: bool) -> list[dict[str, Any]]:
    clauses: list[dict[str, Any]] = []
    for line_number, line in enumerate(lines, start=1):
        if len(line) < min_length:
            continue
        match = CLAUSE_PREFIX_RE.match(line)
        has_normative = contains_any(line, NORMATIVE_KEYWORDS)
        has_acceptance = extract_acceptance_criteria(line) is not None
        has_evidence = any(contains_any(line, keywords) for keywords in EVIDENCE_KEYWORDS.values())
        if not (include_all or match or has_normative or has_acceptance or has_evidence):
            continue
        raw_id = match.group(1).strip() if match else None
        text = match.group(2).strip() if match else line
        clause_id = raw_id.replace(" ", ".").replace(")", "") if raw_id else f"CL-{len(clauses) + 1:03d}"
        flags: list[str] = []
        if not raw_id:
            flags.append("auto_id")
        if not has_normative:
            flags.append("normative_language_not_explicit")
        if has_evidence and "inspection_record" in infer_evidence(text):
            flags.append("evidence_type_unclear")
        if "standart" in text.lower() or "standard" in text.lower():
            flags.append("external_standard_reference")
        criteria = extract_acceptance_criteria(text)
        if not criteria and infer_category(text) in {"test", "tolerance"}:
            flags.append("acceptance_criteria_unclear")
        clauses.append(
            {
                "id": clause_id,
                "source_ref": f"line {line_number}",
                "text": text,
                "category": infer_category(text),
                "evidence_required": infer_evidence(text),
                "acceptance_criteria": criteria,
                "severity": infer_severity(text),
                "flags": flags,
            }
        )
    return clauses


def to_markdown(payload: dict[str, Any]) -> str:
    lines = [
        "# Şartname Madde Adayları",
        "",
        f"- Kaynak: `{payload['metadata']['source']}`",
        f"- Madde sayısı: {len(payload['clauses'])}",
        "",
        "| ID | Kategori | Kanıt | Şiddet | Metin |",
        "|---|---|---|---|---|",
    ]
    for clause in payload["clauses"]:
        evidence = ", ".join(clause["evidence_required"])
        text = clause["text"].replace("|", "\\|")
        lines.append(f"| {clause['id']} | {clause['category']} | {evidence} | {clause['severity']} | {text} |")
    if payload["metadata"]["warnings"]:
        lines.extend(["", "## Uyarılar", ""])
        lines.extend(f"- {warning}" for warning in payload["metadata"]["warnings"])
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", help="TXT, MD, CSV, JSON, PDF veya DOCX kaynak dosya")
    parser.add_argument("--output", help="Çıktı dosyası")
    parser.add_argument("--format", choices=["json", "markdown"], default="json")
    parser.add_argument("--min-length", type=int, default=24)
    parser.add_argument("--include-all", action="store_true", help="Normatif kelime aramadan tüm uzun satırları dahil et")
    args = parser.parse_args()

    source = Path(args.source)
    text, warnings = read_text(source)
    clauses = extract_clauses(normalize_lines(text), args.min_length, args.include_all)
    payload = {
        "metadata": {
            "source": str(source),
            "source_type": source.suffix.lower().lstrip(".") or "unknown",
            "created_at": datetime.now(timezone.utc).isoformat(),
            "extraction_confidence": "low" if warnings else "medium",
            "warnings": warnings,
        },
        "clauses": clauses,
    }
    rendered = to_markdown(payload) if args.format == "markdown" else json.dumps(payload, ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(rendered, encoding="utf-8")
    else:
        print(rendered)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
